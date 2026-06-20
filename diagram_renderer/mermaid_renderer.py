"""Mermaid diagram renderer.

Rendering pipeline for PNG:
  1. Try local ``mmdc`` (Node.js @mermaid-js/mermaid-cli) — best quality.
  2. Fall back to the online ``mermaid.ink`` API — no extra install needed.

Rendering pipeline for PPTX:
  • flowchart / graph  → native PowerPoint shapes (via mermaid_layout + pptx_exporter).
  • all other types    → PNG rendered first, then embedded centred on a blank slide.
"""
from __future__ import annotations

import base64
import shutil
import subprocess
import tempfile
import urllib.error
import urllib.request
from pathlib import Path

from PIL import Image as _PilImage
from pptx import Presentation
from pptx.util import Emu

from .mermaid_layout import is_native_mermaid, mermaid_to_layout
from .pptx_exporter import export_layout_to_pptx
from .style import BG

# ─── Mermaid keyword detection ────────────────────────────────────────────────
_MERMAID_STARTS: tuple[str, ...] = (
    "graph ",      "graph\t",     "graph\n",
    "flowchart ",  "flowchart\t", "flowchart\n",
    "sequencediagram",
    "classdiagram",
    "statediagram",
    "erdiagram",
    "gantt",
    "pie ",        "pie\n",
    "gitgraph",
    "journey",
    "mindmap",
    "timeline",
    "quadrantchart",
    "requirementdiagram",
    "c4context",
    "block-beta",
    "sankey-beta",
    "xychart-beta",
    "architecture-beta",
    "packet-beta",
)

# ─── Slide dimensions (EMU) — matches pptx_exporter ──────────────────────────
_SLIDE_W = 9144000   # 10 inches
_SLIDE_H = 5143500   # 5.625 inches


def is_mermaid(text: str) -> bool:
    """Return True if *text* looks like a Mermaid diagram."""
    stripped = text.strip()
    lower    = stripped.lower()

    # ```mermaid fence → definitely Mermaid
    for line in lower.splitlines():
        s = line.strip()
        if s == "```mermaid":
            return True
        if s:   # first non-empty line reached without finding ```mermaid
            break

    # Keyword at start of diagram text (after optional fence)
    content = _strip_fence(stripped).strip().lower()
    for kw in _MERMAID_STARTS:
        if content.startswith(kw):
            return True
    return False


def _strip_fence(text: str) -> str:
    """Remove optional ```mermaid / ``` wrapper."""
    lines = text.strip().splitlines()
    if lines and lines[0].strip().lower() in ("```mermaid", "```"):
        lines = lines[1:]
    if lines and lines[-1].strip() == "```":
        lines = lines[:-1]
    return "\n".join(lines)


# ─── Public render functions ──────────────────────────────────────────────────

def render_mermaid_to_png(text: str, output_path: Path) -> Path:
    """Render a Mermaid diagram to *output_path* (PNG).

    Tries ``mmdc`` first; falls back to mermaid.ink online API.
    """
    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    diagram = _strip_fence(text)

    if not _try_mmdc(diagram, output_path):
        _render_via_ink(diagram, output_path)

    return output_path


def render_mermaid_to_pptx(text: str, output_path: Path) -> Path:
    """Render a Mermaid diagram to *output_path* (PPTX).

    * flowchart / graph → native editable shapes.
    * all other types   → PNG embedded centred on a blank slide (fallback).
    """
    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    diagram = _strip_fence(text)

    if is_native_mermaid(diagram):
        layout = mermaid_to_layout(diagram)
        return export_layout_to_pptx(layout, output_path)

    # Fallback: render to PNG then embed
    with tempfile.NamedTemporaryFile(suffix=".png", delete=False) as tmp:
        tmp_png = Path(tmp.name)

    try:
        if not _try_mmdc(diagram, tmp_png):
            _render_via_ink(diagram, tmp_png)
        _embed_png_in_pptx(tmp_png, output_path)
    finally:
        tmp_png.unlink(missing_ok=True)

    return output_path


# ─── Internal helpers ─────────────────────────────────────────────────────────

def _try_mmdc(text: str, output_path: Path) -> bool:
    """Attempt rendering with the local ``mmdc`` CLI.  Returns True on success."""
    mmdc = shutil.which("mmdc")
    if mmdc is None:
        return False

    try:
        with tempfile.NamedTemporaryFile(
            mode="w", suffix=".mmd", delete=False, encoding="utf-8"
        ) as f:
            f.write(text)
            tmp_in = Path(f.name)

        result = subprocess.run(
            [mmdc, "-i", str(tmp_in), "-o", str(output_path), "-b", "white"],
            capture_output=True,
            timeout=30,
        )
        tmp_in.unlink(missing_ok=True)
        return result.returncode == 0 and output_path.exists()

    except (FileNotFoundError, subprocess.TimeoutExpired, OSError):
        return False


def _render_via_ink(text: str, output_path: Path) -> None:
    """Render via the mermaid.ink online API (no local Node.js required)."""
    encoded = base64.urlsafe_b64encode(text.encode("utf-8")).decode("ascii")
    url = f"https://mermaid.ink/img/{encoded}?bgColor={BG.lstrip('#')}"

    try:
        req = urllib.request.Request(
            url, headers={"User-Agent": "diagram-ascii2png/1.0"}
        )
        with urllib.request.urlopen(req, timeout=15) as resp:
            output_path.write_bytes(resp.read())
    except urllib.error.URLError as exc:
        raise RuntimeError(
            "Mermaid 렌더링 실패.\n"
            "  • mmdc 설치: npm install -g @mermaid-js/mermaid-cli\n"
            "  • 또는 인터넷 연결을 확인해 주세요 (online fallback: mermaid.ink)\n"
            f"  상세: {exc}"
        ) from exc


def _embed_png_in_pptx(png_path: Path, output_path: Path) -> None:
    """Create a one-slide PPTX with *png_path* centred and aspect-ratio fitted."""
    margin = int(_SLIDE_W * 0.04)   # ~4% margin on each side
    max_w  = _SLIDE_W - margin * 2
    max_h  = _SLIDE_H - margin * 2

    with _PilImage.open(str(png_path)) as img:
        img_w, img_h = img.size

    # Scale to fit while keeping aspect ratio
    scale   = min(max_w / img_w, max_h / img_h)
    final_w = int(img_w * scale)
    final_h = int(img_h * scale)
    left    = (_SLIDE_W - final_w) // 2
    top     = (_SLIDE_H - final_h) // 2

    prs = Presentation()
    prs.slide_width  = Emu(_SLIDE_W)
    prs.slide_height = Emu(_SLIDE_H)

    blank_layout = prs.slide_layouts[6]   # "Blank" layout
    slide = prs.slides.add_slide(blank_layout)
    slide.shapes.add_picture(
        str(png_path),
        left=left, top=top,
        width=final_w, height=final_h,
    )
    prs.save(str(output_path))
